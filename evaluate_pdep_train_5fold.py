#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
 5-Fold Cross-Validation Evaluation of P-Dependent (Parametrized) Fusion Rules
 on the 80% TRAINING folds of a multibiometric score database.
================================================================================

WHAT THIS SCRIPT DOES
---------------------
Given a .mat file containing four 517x517 biometric score matrices (two face
systems and two fingerprint systems), this script:

  1. Splits the subjects into K folds (default 5) using a fixed random seed.
  2. For each fold, forms an 80% TRAIN subset (the other K-1 folds).
  3. OPTIONALLY fits min-max normalization on the TRAIN subset ONLY (no leakage).
     Skip this step with --no-normalize if the matrices are already in [0, 1].
  4. Builds 8 fusion systems (S1..S8) from the score matrices.
  5. Sweeps every (rule, p) combination of the parametrized fusion rules on the
     TRAIN subset, computing EER and Rank-1 @ FAR=1%.
  6. Aggregates results across folds (mean, std, min, max, 95% CI).
  7. Writes a multi-sheet, formatted .xlsx report with full details.

The script evaluates ONLY on the TRAIN folds -- it is the "tuning stage" of a
train/test protocol. The held-out TEST folds are never touched here.

INVALID SCORES
--------------
Invalid / missing comparisons are excluded from normalization and from every
metric. Two conventions are supported, selected with --invalid-value:

    --invalid-value -1     entries equal to -1 are invalid   (default)
    --invalid-value nan    entries that are NaN are invalid

NaN is handled with NaN-safe comparisons throughout (`x != nan` is always True,
so a plain equality test would silently fail to detect it).

USAGE
-----
    python evaluate_pdep_train_5fold.py
    python evaluate_pdep_train_5fold.py --mat-file scores.mat --n-folds 10
    python evaluate_pdep_train_5fold.py --p-step 0.05 --seed 123
    python evaluate_pdep_train_5fold.py --output my_report.xlsx --no-excel

    # Matrices already normalized to [0, 1], invalid entries marked as NaN:
    python evaluate_pdep_train_5fold.py --no-normalize --invalid-value nan

All behaviour is controlled by the CONFIG dataclasses below and/or CLI flags.

REQUIREMENTS
------------
    numpy, pandas, scipy, openpyxl
================================================================================
"""

from __future__ import annotations

import argparse
import logging
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from scipy import stats

# Excel writing is optional -- the script still produces CSVs without it.
try:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover
    OPENPYXL_AVAILABLE = False


# ==============================================================================
#  SECTION 1 -- CONFIGURATION
#  Everything a user would reasonably want to change lives here.
# ==============================================================================

@dataclass
class DataConfig:
    """Where the score data lives and how it is structured."""

    # Path to the MATLAB .mat file holding the score matrices.
    mat_file: str = "870d37a7-673d-49ff-8a98-47b9b519237f.mat"

    # Maps a friendly modality name -> the variable name inside the .mat file.
    # Change the right-hand side if your .mat uses different variable names.
    matrix_keys: Dict[str, str] = field(default_factory=lambda: {
        "Face_C":    "fingToFaceCSCORES",
        "Face_G":    "fingToFaceGSCORES",
        "Finger_LI": "VfingToFaceliSCORES",
        "Finger_RI": "VfingToFaceriSCORES",
    })

    # Number of enrolled subjects (matrices are expected to be n_subjects^2).
    # Set to None to infer automatically from the first matrix loaded.
    n_subjects: Optional[int] = None

    # Sentinel value marking an invalid / missing score. These entries are
    # excluded from normalization and from all metric computations.
    # Use float("nan") if your matrices mark invalid entries with NaN.
    invalid_value: float = -1.0

    # Apply min-max normalization (fitted on TRAIN only) before fusion.
    # Set to False when the matrices are ALREADY normalized to [0, 1] --
    # the fusion rules assume fuzzy-valued inputs in that range.
    normalize: bool = True

    # When normalization is skipped, verify valid scores really do lie in
    # [0, 1] and abort with a clear message if they do not.
    validate_range_if_not_normalized: bool = True


@dataclass
class SplitConfig:
    """How subjects are partitioned into cross-validation folds."""

    n_folds: int = 5           # K in K-fold cross-validation
    seed: int = 42             # RNG seed -- fixes the split for reproducibility
    shuffle: bool = True       # Permute subjects before splitting into folds


@dataclass
class SweepConfig:
    """Which fusion rules and which values of p to evaluate."""

    # Values of the compensation parameter p to sweep.
    p_min: float = 0.0
    p_max: float = 1.0
    p_step: float = 0.1

    # Which rules to include. Must be keys of FUSION_RULES (Section 3).
    # Set to None to use every registered rule.
    rules: Optional[Sequence[str]] = None

    # Which fusion systems to evaluate. Must be keys of SYSTEM_DEFINITIONS.
    # Set to None to use every registered system.
    systems: Optional[Sequence[str]] = None

    def p_values(self) -> List[float]:
        """Return the list of p values to sweep, rounded to avoid FP noise."""
        n_steps = int(round((self.p_max - self.p_min) / self.p_step)) + 1
        return [round(self.p_min + i * self.p_step, 10) for i in range(n_steps)]


@dataclass
class MetricConfig:
    """Parameters of the biometric performance metrics."""

    # Target False Accept Rate for the Rank-1 identification metric.
    far_target: float = 0.01

    # Confidence level used for the CI reported alongside mean/std.
    confidence_level: float = 0.95


@dataclass
class OutputConfig:
    """What gets written, and where."""

    output_xlsx: str = "PDependent_TRAIN_5FoldCV_report.xlsx"
    output_dir: str = "."
    write_excel: bool = True
    write_csv: bool = True

    # Round p values to this many decimals in the report (display only).
    p_decimals: int = 2

    # Colour used for Excel header fills and titles (hex, no leading '#').
    theme_color: str = "0B5394"


@dataclass
class Config:
    """Top-level container bundling all configuration sections."""

    data: DataConfig = field(default_factory=DataConfig)
    split: SplitConfig = field(default_factory=SplitConfig)
    sweep: SweepConfig = field(default_factory=SweepConfig)
    metrics: MetricConfig = field(default_factory=MetricConfig)
    output: OutputConfig = field(default_factory=OutputConfig)
    verbose: bool = True


# ==============================================================================
#  SECTION 2 -- INVALID-ENTRY HANDLING AND SCORE NORMALIZATION
# ==============================================================================

def is_invalid(matrix: np.ndarray, invalid_value: float) -> np.ndarray:
    """
    Boolean mask of invalid entries, safe for both numeric and NaN sentinels.

    A plain `matrix == np.nan` is always False, so NaN sentinels must be tested
    with np.isnan instead. Every other part of the pipeline routes through this
    helper so the two conventions behave identically.
    """
    if isinstance(invalid_value, float) and np.isnan(invalid_value):
        return np.isnan(matrix)
    # NaNs are always treated as invalid, whatever the configured sentinel is.
    return (matrix == invalid_value) | np.isnan(matrix)


def is_valid(matrix: np.ndarray, invalid_value: float) -> np.ndarray:
    """Boolean mask of usable entries -- the complement of `is_invalid`."""
    return ~is_invalid(matrix, invalid_value)


def valid_entries(matrix: np.ndarray, invalid_value: float) -> np.ndarray:
    """Flat array of the usable scores in `matrix`."""
    return matrix[is_valid(matrix, invalid_value)]


class MinMaxNormalizer:
    """
    Min-max normalizer fitted on a TRAIN submatrix and applied to any submatrix.

    Fitting on train data only is what keeps the protocol leakage-free: the
    held-out subjects never influence the scale their own scores are mapped to.

    Invalid entries are excluded from the fit and passed through untouched by
    `transform`, keeping whatever sentinel the caller configured.
    """

    def __init__(self, invalid_value: float = -1.0) -> None:
        self.invalid_value = invalid_value
        self.vmin: Optional[float] = None
        self.vmax: Optional[float] = None

    def fit(self, train_matrix: np.ndarray) -> "MinMaxNormalizer":
        """Learn min and max from the valid entries of `train_matrix`."""
        valid = valid_entries(train_matrix, self.invalid_value)
        if valid.size == 0:
            raise ValueError("No valid entries found to fit the normalizer.")
        self.vmin = float(valid.min())
        self.vmax = float(valid.max())
        if np.isclose(self.vmax, self.vmin):
            raise ValueError(
                f"Degenerate score range: min == max == {self.vmin}. "
                "Cannot min-max normalize."
            )
        return self

    def transform(self, matrix: np.ndarray) -> np.ndarray:
        """Rescale valid entries to [0, 1]; leave invalid entries unchanged."""
        if self.vmin is None or self.vmax is None:
            raise RuntimeError("Normalizer must be fitted before transform().")
        out = np.full(matrix.shape, self.invalid_value, dtype=float)
        mask = is_valid(matrix, self.invalid_value)
        out[mask] = (matrix[mask] - self.vmin) / (self.vmax - self.vmin)
        return out

    def fit_transform(self, train_matrix: np.ndarray) -> np.ndarray:
        return self.fit(train_matrix).transform(train_matrix)


class PassthroughNormalizer:
    """
    No-op stand-in used when `DataConfig.normalize` is False.

    Exposes the same interface as MinMaxNormalizer so the pipeline does not
    need to branch. Optionally checks that the incoming scores really are in
    [0, 1], since the fusion rules assume fuzzy-valued inputs.
    """

    def __init__(self, invalid_value: float = -1.0,
                 validate_range: bool = True,
                 label: str = "") -> None:
        self.invalid_value = invalid_value
        self.validate_range = validate_range
        self.label = label
        self.vmin: Optional[float] = None
        self.vmax: Optional[float] = None

    def fit(self, train_matrix: np.ndarray) -> "PassthroughNormalizer":
        valid = valid_entries(train_matrix, self.invalid_value)
        if valid.size == 0:
            raise ValueError("No valid entries found in the score matrix.")
        self.vmin = float(valid.min())
        self.vmax = float(valid.max())

        if self.validate_range and (self.vmin < -1e-9 or self.vmax > 1.0 + 1e-9):
            where = f" for '{self.label}'" if self.label else ""
            raise ValueError(
                f"Normalization is disabled but scores{where} fall outside "
                f"[0, 1] (observed range [{self.vmin:.6g}, {self.vmax:.6g}]). "
                "The fusion rules require inputs in [0, 1]. Either enable "
                "normalization, or pass already-normalized matrices."
            )
        return self

    def transform(self, matrix: np.ndarray) -> np.ndarray:
        """Return the matrix unchanged (invalid entries keep their sentinel)."""
        return matrix.astype(float, copy=True)

    def fit_transform(self, train_matrix: np.ndarray) -> np.ndarray:
        return self.fit(train_matrix).transform(train_matrix)


def make_normalizer(cfg: DataConfig, label: str = ""):
    """Return the normalizer implied by the configuration."""
    if cfg.normalize:
        return MinMaxNormalizer(cfg.invalid_value)
    return PassthroughNormalizer(
        cfg.invalid_value,
        validate_range=cfg.validate_range_if_not_normalized,
        label=label,
    )


# ==============================================================================
#  SECTION 3 -- FUSION RULES
#
#  Each rule is a function f(x, y, p) -> fused score, operating element-wise on
#  arrays already normalized to [0, 1]. Register new rules by adding them to the
#  FUSION_RULES dict; they become available to the sweep automatically.
# ==============================================================================

def _safe_pow(base: np.ndarray, exponent: float) -> np.ndarray:
    """
    Element-wise power that handles the 0^0 edge case explicitly.

    numpy returns 1.0 for 0**0, which is the convention we want, but negative
    bases with fractional exponents produce NaN -- so we clip to >= 0 first.
    """
    base_clipped = np.clip(base, 0.0, None)
    with np.errstate(invalid="ignore", divide="ignore"):
        result = np.power(base_clipped, exponent,
                          where=(base_clipped > 0),
                          out=np.zeros_like(base_clipped))
    zero_result = 1.0 if exponent == 0 else 0.0
    return np.where(base_clipped == 0, zero_result, result)


def rule_zimmermann1(x: np.ndarray, y: np.ndarray, p: float) -> np.ndarray:
    """F = (xy)^(1-p) * (x + y - xy)^p"""
    product = x * y
    prob_sum = x + y - product
    return _safe_pow(product, 1.0 - p) * _safe_pow(prob_sum, p)


def rule_zimmermann2(x: np.ndarray, y: np.ndarray, p: float) -> np.ndarray:
    """F = min(x,y)(1-p) + max(x,y) p"""
    return np.minimum(x, y) * (1 - p) + np.maximum(x, y) * p


def rule_luhandjula(x: np.ndarray, y: np.ndarray, p: float) -> np.ndarray:
    """F = min(x,y)(1-p) + min(1, x+y) p"""
    return np.minimum(x, y) * (1 - p) + np.minimum(1.0, x + y) * p


def rule_werners1(x: np.ndarray, y: np.ndarray, p: float) -> np.ndarray:
    """F = min(x,y)(1-p) + ((x+y)/2) p"""
    return np.minimum(x, y) * (1 - p) + ((x + y) / 2.0) * p


def rule_werners2(x: np.ndarray, y: np.ndarray, p: float) -> np.ndarray:
    """F = max(x,y)(1-p) + ((x+y)/2) p"""
    return np.maximum(x, y) * (1 - p) + ((x + y) / 2.0) * p


def rule_sales1(x: np.ndarray, y: np.ndarray, p: float) -> np.ndarray:
    """F = max(0, x+y-1)(1-p) + min(x,y) p"""
    return np.maximum(0.0, x + y - 1.0) * (1 - p) + np.minimum(x, y) * p


def rule_sales2(x: np.ndarray, y: np.ndarray, p: float) -> np.ndarray:
    """F = min(1, x+y)(1-p) + max(x,y) p"""
    return np.minimum(1.0, x + y) * (1 - p) + np.maximum(x, y) * p


def rule_mizumoto1(x: np.ndarray, y: np.ndarray, p: float) -> np.ndarray:
    """F = min(x,y)(1-p) + sqrt(xy) p"""
    return np.minimum(x, y) * (1 - p) + np.sqrt(np.clip(x * y, 0.0, None)) * p


def rule_mizumoto2(x: np.ndarray, y: np.ndarray, p: float) -> np.ndarray:
    """F = ((x+y)/2)(1-p) + sqrt(xy) p"""
    return ((x + y) / 2.0) * (1 - p) + np.sqrt(np.clip(x * y, 0.0, None)) * p


def rule_r5p(x: np.ndarray, y: np.ndarray, p: float) -> np.ndarray:
    """Proposed rule: F = p(x+y-xy) + (1-p)(1-(x+y-xy)) min(x,y)"""
    prob_sum = x + y - x * y
    return p * prob_sum + (1 - p) * (1 - prob_sum) * np.minimum(x, y)


# Registry: rule name -> implementation. Add your own rules here.
FUSION_RULES: Dict[str, Callable[[np.ndarray, np.ndarray, float], np.ndarray]] = {
    "Zimmermann1": rule_zimmermann1,
    "Zimmermann2": rule_zimmermann2,
    "Luhandjula":  rule_luhandjula,
    "Werners1":    rule_werners1,
    "Werners2":    rule_werners2,
    "Sales1":      rule_sales1,
    "Sales2":      rule_sales2,
    "Mizumoto1":   rule_mizumoto1,
    "Mizumoto2":   rule_mizumoto2,
    "R5p":         rule_r5p,
}

# Preferred display order for rules in reports (unlisted rules are appended).
RULE_DISPLAY_ORDER: List[str] = list(FUSION_RULES.keys())


def apply_fusion_rule(rule_name: str,
                      x: np.ndarray,
                      y: np.ndarray,
                      p: float,
                      invalid_value: float = -1.0) -> np.ndarray:
    """
    Fuse two normalized score matrices with the named rule at parameter p.

    Entries invalid in EITHER input stay invalid in the output. Valid outputs
    are clipped to [0, 1] so downstream rules (in cascaded systems) always
    receive well-formed fuzzy inputs.
    """
    if rule_name not in FUSION_RULES:
        raise KeyError(f"Unknown fusion rule '{rule_name}'. "
                       f"Available: {sorted(FUSION_RULES)}")

    valid = is_valid(x, invalid_value) & is_valid(y, invalid_value)
    # Substitute 0 in invalid positions so arithmetic never sees the sentinel
    # (critical for NaN, which would otherwise propagate through every term).
    x_safe = np.where(valid, x, 0.0)
    y_safe = np.where(valid, y, 0.0)

    with np.errstate(invalid="ignore", divide="ignore"):
        fused = FUSION_RULES[rule_name](x_safe, y_safe, p)

    fused = np.clip(fused, 0.0, 1.0)
    return np.where(valid, fused, invalid_value)


# ==============================================================================
#  SECTION 4 -- FUSION SYSTEM DEFINITIONS
#
#  A system is a recipe for combining modalities. Single-stage systems fuse two
#  modalities; cascaded systems fuse two modalities, then fuse the result with a
#  third, reusing the same rule and the same p at both stages.
# ==============================================================================

@dataclass(frozen=True)
class SystemDefinition:
    """Describes how one fusion system is assembled from base modalities."""

    name: str
    description: str
    stage1: Tuple[str, str]          # the two modalities fused first
    stage2_modality: Optional[str] = None  # if set, fuse stage-1 output with this

    @property
    def is_cascaded(self) -> bool:
        return self.stage2_modality is not None


SYSTEM_DEFINITIONS: Dict[str, SystemDefinition] = {
    "S1": SystemDefinition("S1", "Face_C + Face_G",
                           ("Face_C", "Face_G")),
    "S2": SystemDefinition("S2", "Face_C + Finger_LI",
                           ("Face_C", "Finger_LI")),
    "S3": SystemDefinition("S3", "Face_C + Finger_RI",
                           ("Face_C", "Finger_RI")),
    "S4": SystemDefinition("S4", "Face_G + Finger_LI",
                           ("Face_G", "Finger_LI")),
    "S5": SystemDefinition("S5", "Face_G + Finger_RI",
                           ("Face_G", "Finger_RI")),
    "S6": SystemDefinition("S6", "Finger_LI + Finger_RI",
                           ("Finger_LI", "Finger_RI")),
    "S7": SystemDefinition("S7", "(Finger_LI + Finger_RI) -> Face_C",
                           ("Finger_LI", "Finger_RI"), "Face_C"),
    "S8": SystemDefinition("S8", "(Finger_LI + Finger_RI) -> Face_G",
                           ("Finger_LI", "Finger_RI"), "Face_G"),
}


def build_system_matrix(system: SystemDefinition,
                        normalized: Dict[str, np.ndarray],
                        rule_name: str,
                        p: float,
                        invalid_value: float = -1.0) -> np.ndarray:
    """Produce the fused score matrix for one system at one (rule, p)."""
    mod_a, mod_b = system.stage1
    fused = apply_fusion_rule(rule_name, normalized[mod_a], normalized[mod_b],
                              p, invalid_value)
    if system.is_cascaded:
        fused = apply_fusion_rule(rule_name, fused,
                                  normalized[system.stage2_modality],
                                  p, invalid_value)
    return fused


# ==============================================================================
#  SECTION 5 -- BIOMETRIC PERFORMANCE METRICS
# ==============================================================================

def split_genuine_impostor(score_matrix: np.ndarray,
                           invalid_value: float = -1.0
                           ) -> Tuple[np.ndarray, np.ndarray]:
    """
    Separate a square score matrix into genuine and impostor score vectors.

    Convention: row i and column i refer to the same subject, so the diagonal
    holds genuine comparisons and everything off-diagonal is an impostor
    comparison. Invalid entries are dropped from both vectors.
    """
    n = score_matrix.shape[0]
    diagonal = np.diag(score_matrix)
    genuine = diagonal[is_valid(diagonal, invalid_value)]

    off_diagonal = score_matrix[~np.eye(n, dtype=bool)]
    impostor = off_diagonal[is_valid(off_diagonal, invalid_value)]

    return genuine, impostor


def compute_eer(genuine: np.ndarray, impostor: np.ndarray) -> float:
    """
    Equal Error Rate: the error rate where FAR and FRR cross.

    Implemented by sweeping every observed score as a candidate threshold and
    taking the point where |FAR - FRR| is smallest. Uses sorted arrays plus
    binary search, so it is O(m log m) rather than O(m * n_thresholds).
    """
    if genuine.size == 0 or impostor.size == 0:
        return float("nan")

    genuine_sorted = np.sort(genuine)
    impostor_sorted = np.sort(impostor)
    thresholds = np.unique(np.concatenate([genuine_sorted, impostor_sorted]))

    # FAR(t) = fraction of impostor scores >= t
    far = 1.0 - np.searchsorted(impostor_sorted, thresholds, side="left") / impostor_sorted.size
    # FRR(t) = fraction of genuine scores < t
    frr = np.searchsorted(genuine_sorted, thresholds, side="left") / genuine_sorted.size

    crossing = int(np.argmin(np.abs(far - frr)))
    return float((far[crossing] + frr[crossing]) / 2.0)


def compute_rank1_at_far(score_matrix: np.ndarray,
                         far_target: float = 0.01,
                         invalid_value: float = -1.0) -> float:
    """
    Open-set Rank-1 identification rate at a fixed False Accept Rate.

    A probe counts as correctly identified only if BOTH hold:
      (a) its highest-scoring gallery entry is the correct subject, and
      (b) that score clears a threshold set so that `far_target` of impostor
          scores would be accepted.
    """
    n = score_matrix.shape[0]
    diagonal = np.diag(score_matrix)

    _, impostor = split_genuine_impostor(score_matrix, invalid_value)
    if impostor.size == 0:
        return float("nan")

    threshold = np.percentile(impostor, 100.0 * (1.0 - far_target))

    # Invalid cells must never win the argmax -> push them to -inf.
    # (NaN would poison both max and argmax, so this substitution is required.)
    searchable = np.where(is_invalid(score_matrix, invalid_value), -np.inf, score_matrix)
    top_scores = searchable.max(axis=1)
    top_indices = searchable.argmax(axis=1)

    has_valid_mate = is_valid(diagonal, invalid_value)
    is_correct = top_indices == np.arange(n)
    is_accepted = top_scores >= threshold

    hits = is_correct & is_accepted & has_valid_mate
    n_probes = int(has_valid_mate.sum())
    return float(hits.sum() / n_probes) if n_probes else float("nan")


# ==============================================================================
#  SECTION 6 -- CROSS-VALIDATION SPLITTING
# ==============================================================================

def make_folds(n_subjects: int, cfg: SplitConfig) -> List[np.ndarray]:
    """
    Partition subject indices into `cfg.n_folds` disjoint, roughly equal folds.

    Returns a list of arrays of subject indices -- one array per fold. Each fold
    serves in turn as the held-out portion; the union of the others is TRAIN.
    """
    indices = np.arange(n_subjects)
    if cfg.shuffle:
        rng = np.random.default_rng(cfg.seed)
        indices = rng.permutation(indices)
    return [np.sort(fold) for fold in np.array_split(indices, cfg.n_folds)]


def train_indices_for_fold(folds: Sequence[np.ndarray],
                           held_out_fold: int,
                           n_subjects: int) -> np.ndarray:
    """Return the sorted TRAIN subject indices when `held_out_fold` is excluded."""
    return np.sort(np.setdiff1d(np.arange(n_subjects), folds[held_out_fold]))


def submatrix(matrix: np.ndarray, subject_ids: np.ndarray) -> np.ndarray:
    """Extract the square submatrix for a subset of subjects (rows AND cols)."""
    return matrix[np.ix_(subject_ids, subject_ids)]


# ==============================================================================
#  SECTION 7 -- THE EVALUATION PIPELINE
# ==============================================================================

class PDependentTrainEvaluator:
    """
    Runs the full (rule x p x system x fold) sweep on the TRAIN folds.

    Typical use:
        evaluator = PDependentTrainEvaluator(config)
        evaluator.load_data()
        results = evaluator.run()
        tables = evaluator.summarize()
    """

    def __init__(self, config: Config) -> None:
        self.cfg = config
        self.log = logging.getLogger(self.__class__.__name__)

        self.raw_matrices: Dict[str, np.ndarray] = {}
        self.n_subjects: int = 0
        self.folds: List[np.ndarray] = []

        self.raw_results: Optional[pd.DataFrame] = None
        self.tables: Dict[str, pd.DataFrame] = {}

    # ---------------------------------------------------------------- loading

    def load_data(self) -> "PDependentTrainEvaluator":
        """Read the .mat file and validate the score matrices."""
        from scipy.io import loadmat  # imported lazily: only needed here

        path = Path(self.cfg.data.mat_file)
        if not path.exists():
            raise FileNotFoundError(f"Score file not found: {path.resolve()}")

        self.log.info("Loading score matrices from %s", path)
        contents = loadmat(str(path))

        for modality, key in self.cfg.data.matrix_keys.items():
            if key not in contents:
                raise KeyError(
                    f"Variable '{key}' (for modality '{modality}') not found in "
                    f"{path.name}. Available: "
                    f"{[k for k in contents if not k.startswith('__')]}"
                )
            matrix = np.asarray(contents[key], dtype=float)
            if matrix.ndim != 2 or matrix.shape[0] != matrix.shape[1]:
                raise ValueError(
                    f"Matrix '{key}' must be square, got shape {matrix.shape}."
                )
            self.raw_matrices[modality] = matrix

        shapes = {m: s.shape[0] for m, s in self.raw_matrices.items()}
        if len(set(shapes.values())) != 1:
            raise ValueError(f"All matrices must share a size; got {shapes}")

        inferred = next(iter(shapes.values()))
        configured = self.cfg.data.n_subjects
        if configured is not None and configured != inferred:
            raise ValueError(
                f"Config says n_subjects={configured} but matrices are {inferred}x{inferred}."
            )
        self.n_subjects = inferred

        sentinel = ("NaN" if np.isnan(self.cfg.data.invalid_value)
                    else self.cfg.data.invalid_value)
        self.log.info("Loaded %d modalities, %d subjects each. "
                      "Invalid sentinel: %s. Normalization: %s.",
                      len(self.raw_matrices), self.n_subjects, sentinel,
                      "min-max (fitted on TRAIN)" if self.cfg.data.normalize
                      else "DISABLED (inputs assumed pre-normalized)")

        for modality, matrix in self.raw_matrices.items():
            invalid_mask = is_invalid(matrix, self.cfg.data.invalid_value)
            valid = matrix[~invalid_mask]
            self.log.info("  %-10s shape=%s  invalid=%d  range=[%.4g, %.4g]",
                          modality, matrix.shape, int(invalid_mask.sum()),
                          valid.min() if valid.size else float("nan"),
                          valid.max() if valid.size else float("nan"))
        return self

    # ------------------------------------------------------------------- sweep

    def _active_rules(self) -> List[str]:
        requested = self.cfg.sweep.rules
        if requested is None:
            return list(FUSION_RULES.keys())
        unknown = set(requested) - set(FUSION_RULES)
        if unknown:
            raise KeyError(f"Unknown rules requested: {sorted(unknown)}")
        return list(requested)

    def _active_systems(self) -> List[SystemDefinition]:
        requested = self.cfg.sweep.systems
        names = list(SYSTEM_DEFINITIONS) if requested is None else list(requested)
        unknown = set(names) - set(SYSTEM_DEFINITIONS)
        if unknown:
            raise KeyError(f"Unknown systems requested: {sorted(unknown)}")
        return [SYSTEM_DEFINITIONS[n] for n in names]

    def _prepare_train_fold(self, train_ids: np.ndarray) -> Dict[str, np.ndarray]:
        """
        Extract and (optionally) normalize one fold's TRAIN submatrices.

        With normalization enabled, min-max parameters are fitted on this
        fold's TRAIN block only -- the held-out subjects never influence it.
        With normalization disabled, the submatrices are returned as-is after
        an optional range check.
        """
        prepared: Dict[str, np.ndarray] = {}
        for modality, matrix in self.raw_matrices.items():
            train_block = submatrix(matrix, train_ids)
            normalizer = make_normalizer(self.cfg.data, label=modality)
            prepared[modality] = normalizer.fit_transform(train_block)
        return prepared

    def run(self) -> pd.DataFrame:
        """Execute the full sweep and return the long-format results table."""
        if not self.raw_matrices:
            raise RuntimeError("Call load_data() before run().")

        rules = self._active_rules()
        systems = self._active_systems()
        p_values = self.cfg.sweep.p_values()
        self.folds = make_folds(self.n_subjects, self.cfg.split)

        total = len(self.folds) * len(systems) * len(rules) * len(p_values)
        self.log.info("Sweeping %d folds x %d systems x %d rules x %d p-values "
                      "= %d evaluations.",
                      len(self.folds), len(systems), len(rules),
                      len(p_values), total)

        records: List[dict] = []
        start = time.perf_counter()

        for fold_idx in range(len(self.folds)):
            train_ids = train_indices_for_fold(self.folds, fold_idx, self.n_subjects)
            normalized = self._prepare_train_fold(train_ids)
            fold_start = time.perf_counter()

            for system in systems:
                for rule_name in rules:
                    for p in p_values:
                        fused = build_system_matrix(
                            system, normalized, rule_name, p,
                            self.cfg.data.invalid_value,
                        )
                        genuine, impostor = split_genuine_impostor(
                            fused, self.cfg.data.invalid_value
                        )
                        records.append({
                            "Fold": fold_idx + 1,
                            "System": system.name,
                            "System_Description": system.description,
                            "Rule": rule_name,
                            "p": round(p, self.cfg.output.p_decimals),
                            "Train_EER": compute_eer(genuine, impostor),
                            "Train_Rank1": compute_rank1_at_far(
                                fused,
                                self.cfg.metrics.far_target,
                                self.cfg.data.invalid_value,
                            ),
                            "N_Train_Subjects": len(train_ids),
                            "N_Genuine": genuine.size,
                            "N_Impostor": impostor.size,
                        })

            self.log.info("  Fold %d/%d done (%d train subjects) in %.1fs",
                          fold_idx + 1, len(self.folds), len(train_ids),
                          time.perf_counter() - fold_start)

        self.raw_results = pd.DataFrame.from_records(records)
        self.log.info("Sweep finished: %d rows in %.1fs",
                      len(self.raw_results), time.perf_counter() - start)
        return self.raw_results

    # --------------------------------------------------------------- summaries

    def _ci_half_width(self, values: pd.Series) -> float:
        """Half-width of the confidence interval for a small sample mean."""
        n = values.count()
        if n < 2:
            return 0.0
        alpha = 1.0 - self.cfg.metrics.confidence_level
        t_crit = stats.t.ppf(1.0 - alpha / 2.0, df=n - 1)
        return float(t_crit * values.std(ddof=1) / np.sqrt(n))

    def _describe(self, values: pd.Series, prefix: str) -> Dict[str, float]:
        """Standard statistic block: mean, std, min, max, CI half-width."""
        return {
            f"{prefix}_mean": float(values.mean()),
            f"{prefix}_std":  float(values.std(ddof=1)) if values.count() > 1 else 0.0,
            f"{prefix}_min":  float(values.min()),
            f"{prefix}_max":  float(values.max()),
            f"{prefix}_CI95": self._ci_half_width(values),
        }

    def summarize(self) -> Dict[str, pd.DataFrame]:
        """
        Build every reporting table from the raw sweep results.

        Returns a dict of sheet-name -> DataFrame, ready to be written to Excel.
        """
        if self.raw_results is None:
            raise RuntimeError("Call run() before summarize().")

        df = self.raw_results
        rule_order = [r for r in RULE_DISPLAY_ORDER if r in set(df["Rule"])]
        system_order = [s for s in SYSTEM_DEFINITIONS if s in set(df["System"])]

        def ordered(frame: pd.DataFrame) -> pd.DataFrame:
            """Sort a frame by the canonical system then rule order."""
            out = frame.copy()
            out["System"] = pd.Categorical(out["System"], system_order, ordered=True)
            if "Rule" in out.columns:
                out["Rule"] = pd.Categorical(out["Rule"], rule_order, ordered=True)
                return out.sort_values(["System", "Rule"]).reset_index(drop=True)
            return out.sort_values("System").reset_index(drop=True)

        tables: Dict[str, pd.DataFrame] = {}

        # --- Fold composition ------------------------------------------------
        tables["Fold_Info"] = pd.DataFrame([
            {
                "Fold": i + 1,
                "N_Test_Subjects": len(fold),
                "N_Train_Subjects": self.n_subjects - len(fold),
                "Train_Percent": round(100 * (self.n_subjects - len(fold)) / self.n_subjects, 2),
            }
            for i, fold in enumerate(self.folds)
        ])

        # --- Configuration audit trail --------------------------------------
        tables["Run_Configuration"] = pd.DataFrame([
            {"Setting": "Score file",        "Value": self.cfg.data.mat_file},
            {"Setting": "Subjects",          "Value": self.n_subjects},
            {"Setting": "Folds",             "Value": self.cfg.split.n_folds},
            {"Setting": "Random seed",       "Value": self.cfg.split.seed},
            {"Setting": "Normalization",     "Value": ("min-max, fitted on TRAIN fold only"
                                                       if self.cfg.data.normalize
                                                       else "DISABLED - inputs assumed already in [0, 1]")},
            {"Setting": "p values",          "Value": ", ".join(str(round(p, self.cfg.output.p_decimals))
                                                                 for p in self.cfg.sweep.p_values())},
            {"Setting": "Rules",             "Value": ", ".join(rule_order)},
            {"Setting": "Systems",           "Value": ", ".join(system_order)},
            {"Setting": "FAR target",        "Value": self.cfg.metrics.far_target},
            {"Setting": "Confidence level",  "Value": self.cfg.metrics.confidence_level},
            {"Setting": "Invalid sentinel",  "Value": ("NaN" if np.isnan(self.cfg.data.invalid_value)
                                                       else self.cfg.data.invalid_value)},
            {"Setting": "Total evaluations", "Value": len(self.raw_results)},
        ])

        # --- Aggregate across folds FIRST, for every (System, Rule, p) --------
        # This ordering matters. Selecting p inside each fold and then averaging
        # would mix several different configurations into one "mean", giving an
        # optimistically biased number that no single deployable system attains.
        # Instead we average each (System, Rule, p) across folds, and only then
        # select -- so every reported winner is ONE fixed (rule, p) pair.
        per_config_rows: List[dict] = []
        for (system, rule, p), group in df.groupby(["System", "Rule", "p"]):
            row = {"System": system, "Rule": rule, "p": p, "N_Folds": len(group)}
            row.update(self._describe(group["Train_EER"], "Train_EER"))
            row.update(self._describe(group["Train_Rank1"], "Train_Rank1"))
            per_config_rows.append(row)
        per_config = pd.DataFrame(per_config_rows)
        per_config["System"] = pd.Categorical(per_config["System"], system_order, ordered=True)
        per_config["Rule"] = pd.Categorical(per_config["Rule"], rule_order, ordered=True)
        per_config = per_config.sort_values(["System", "Rule", "p"]).reset_index(drop=True)
        tables["Stats_per_Config"] = per_config

        # --- Select the best p per (System, Rule), on MEAN EER across folds ---
        best_eer_idx = per_config.groupby(["System", "Rule"], observed=True)["Train_EER_mean"].idxmin()
        stats_by_eer = per_config.loc[best_eer_idx].reset_index(drop=True)
        stats_by_eer = stats_by_eer.rename(columns={"p": "p_selected"})
        tables["Stats_by_Rule_EER"] = ordered(stats_by_eer)

        # Same, but selecting the p that maximizes MEAN Rank-1 across folds.
        best_r1_idx = per_config.groupby(["System", "Rule"], observed=True)["Train_Rank1_mean"].idxmax()
        stats_by_r1 = per_config.loc[best_r1_idx].reset_index(drop=True)
        stats_by_r1 = stats_by_r1.rename(columns={"p": "p_selected"})
        tables["Stats_by_Rule_Rank1"] = ordered(stats_by_r1)

        # --- Per-fold detail for the selected configurations ------------------
        # Shows the five individual fold scores behind each reported mean, using
        # the SAME fixed p in every fold (not a per-fold re-optimized p).
        eer_choice = stats_by_eer[["System", "Rule", "p_selected"]].rename(
            columns={"p_selected": "p"})
        tables["FoldDetail_by_EER"] = ordered(
            df.merge(eer_choice, on=["System", "Rule", "p"], how="inner")
              [["Fold", "System", "Rule", "p", "Train_EER", "Train_Rank1"]]
        ).sort_values(["System", "Rule", "Fold"]).reset_index(drop=True)

        r1_choice = stats_by_r1[["System", "Rule", "p_selected"]].rename(
            columns={"p_selected": "p"})
        tables["FoldDetail_by_Rank1"] = ordered(
            df.merge(r1_choice, on=["System", "Rule", "p"], how="inner")
              [["Fold", "System", "Rule", "p", "Train_EER", "Train_Rank1"]]
        ).sort_values(["System", "Rule", "Fold"]).reset_index(drop=True)

        # --- Wide pivots (quick visual scan) ---------------------------------
        tables["Pivot_Mean_EER"] = (
            stats_by_eer.pivot(index="System", columns="Rule", values="Train_EER_mean")
            .reindex(index=system_order, columns=rule_order)
            .reset_index()
        )
        tables["Pivot_Std_EER"] = (
            stats_by_eer.pivot(index="System", columns="Rule", values="Train_EER_std")
            .reindex(index=system_order, columns=rule_order)
            .reset_index()
        )
        tables["Pivot_Mean_Rank1"] = (
            stats_by_r1.pivot(index="System", columns="Rule", values="Train_Rank1_mean")
            .reindex(index=system_order, columns=rule_order)
            .reset_index()
        )
        tables["Pivot_Std_Rank1"] = (
            stats_by_r1.pivot(index="System", columns="Rule", values="Train_Rank1_std")
            .reindex(index=system_order, columns=rule_order)
            .reset_index()
        )
        tables["Pivot_Selected_p"] = (
            stats_by_eer.pivot(index="System", columns="Rule", values="p_selected")
            .reindex(index=system_order, columns=rule_order)
            .reset_index()
        )

        # --- Ranking of rules within each system ------------------------------
        ranked = stats_by_eer.copy()
        ranked["Rank_in_System"] = (
            ranked.groupby("System", observed=True)["Train_EER_mean"]
            .rank(method="min").astype(int)
        )
        tables["Rules_Ranked_per_System"] = (
            ranked.sort_values(["System", "Train_EER_mean"]).reset_index(drop=True)
        )

        # --- Overall winner per system ----------------------------------------
        # Selected on MEAN EER across folds -- one fixed (rule, p) per system.
        tables["Overall_Best_per_System"] = ordered(
            pd.DataFrame([
                group.loc[group["Train_EER_mean"].idxmin()]
                for _, group in stats_by_eer.groupby("System", observed=True)
            ])
        )

        # Same, selected on MEAN Rank-1 instead.
        tables["Overall_Best_by_Rank1"] = ordered(
            pd.DataFrame([
                group.loc[group["Train_Rank1_mean"].idxmax()]
                for _, group in stats_by_r1.groupby("System", observed=True)
            ])
        )

        # --- Raw sweep --------------------------------------------------------
        tables["Full_Raw_Sweep"] = df

        self.tables = tables
        return tables


# ==============================================================================
#  SECTION 8 -- EXCEL REPORT WRITER
# ==============================================================================

class ExcelReportWriter:
    """Writes the summary tables into a formatted, multi-sheet workbook."""

    # Columns matching these substrings are formatted as percentages.
    PERCENT_HINTS = ("EER", "Rank1")
    # Columns matching these are formatted as plain decimals, not percentages.
    DECIMAL_HINTS = ("p_selected", "Train_Percent")

    # Human-readable blurb shown at the top of each sheet.
    SHEET_TITLES = {
        "Run_Configuration":       "Run configuration and audit trail",
        "Fold_Info":               "Cross-validation fold composition",
        "Overall_Best_per_System": "Winning rule per system - lowest MEAN TRAIN EER across folds",
        "Overall_Best_by_Rank1":   "Winning rule per system - highest MEAN TRAIN Rank-1 across folds",
        "Pivot_Mean_EER":          "Mean TRAIN EER across folds - System x Rule",
        "Pivot_Std_EER":           "Std-dev of TRAIN EER across folds - System x Rule",
        "Pivot_Mean_Rank1":        "Mean TRAIN Rank-1 @ FAR across folds - System x Rule",
        "Pivot_Std_Rank1":         "Std-dev of TRAIN Rank-1 across folds - System x Rule",
        "Pivot_Selected_p":        "Selected p (lowest mean EER) - System x Rule",
        "Stats_by_Rule_EER":       "Per System x Rule at its EER-optimal p: mean, std, min, max, CI95",
        "Stats_by_Rule_Rank1":     "Per System x Rule at its Rank1-optimal p: mean, std, min, max, CI95",
        "Stats_per_Config":        "Every (System, Rule, p) averaged across folds - the base table",
        "Rules_Ranked_per_System": "Rules ranked within each system by mean TRAIN EER",
        "FoldDetail_by_EER":       "Individual fold scores behind each EER-selected configuration",
        "FoldDetail_by_Rank1":     "Individual fold scores behind each Rank1-selected configuration",
        "Full_Raw_Sweep":          "Complete raw sweep - every fold, system, rule and p",
    }

    def __init__(self, config: Config) -> None:
        self.cfg = config
        self.log = logging.getLogger(self.__class__.__name__)
        self.font = "Arial"

    # ------------------------------------------------------------- formatting

    def _styles(self) -> dict:
        color = self.cfg.output.theme_color
        thin = Side(style="thin", color="D9D9D9")
        return {
            "header_fill": PatternFill("solid", fgColor=color),
            "header_font": Font(name=self.font, bold=True, color="FFFFFF", size=10),
            "title_font":  Font(name=self.font, bold=True, size=13, color=color),
            "note_font":   Font(name=self.font, italic=True, size=9, color="555555"),
            "cell_font":   Font(name=self.font, size=10),
            "border":      Border(left=thin, right=thin, top=thin, bottom=thin),
        }

    def _number_format(self, column: str) -> Optional[str]:
        if any(hint in column for hint in self.DECIMAL_HINTS):
            return "0.00"
        if any(hint in column for hint in self.PERCENT_HINTS):
            return "0.00%"
        return None

    def _write_table(self, worksheet, frame: pd.DataFrame, styles: dict,
                     title: Optional[str], note: Optional[str]) -> None:
        row = 1
        if title:
            worksheet.cell(row=row, column=1, value=title).font = styles["title_font"]
            row += 1
        if note:
            worksheet.cell(row=row, column=1, value=note).font = styles["note_font"]
            row += 1
        row += 1

        header_row = row
        for col_idx, column in enumerate(frame.columns, start=1):
            cell = worksheet.cell(row=header_row, column=col_idx, value=str(column))
            cell.fill = styles["header_fill"]
            cell.font = styles["header_font"]
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
            cell.border = styles["border"]

        formats = {c: self._number_format(str(c)) for c in frame.columns}
        for offset, record in enumerate(frame.itertuples(index=False), start=1):
            for col_idx, value in enumerate(record, start=1):
                column = frame.columns[col_idx - 1]
                if isinstance(value, (np.integer,)):
                    value = int(value)
                elif isinstance(value, (np.floating,)):
                    value = float(value)
                elif not isinstance(value, (int, float, str)) and value is not None:
                    value = str(value)
                cell = worksheet.cell(row=header_row + offset, column=col_idx, value=value)
                cell.font = styles["cell_font"]
                cell.border = styles["border"]
                if formats[column]:
                    cell.number_format = formats[column]

        worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=1)
        self._autofit(worksheet, len(frame.columns))

    @staticmethod
    def _autofit(worksheet, n_columns: int, minimum: int = 10, maximum: int = 40) -> None:
        for col_idx in range(1, n_columns + 1):
            letter = get_column_letter(col_idx)
            longest = max(
                (len(str(cell.value)) for cell in worksheet[letter] if cell.value is not None),
                default=minimum,
            )
            worksheet.column_dimensions[letter].width = min(max(longest + 2, minimum), maximum)

    def _add_heatmap(self, worksheet, frame: pd.DataFrame,
                     header_row: int, lower_is_better: bool) -> None:
        """Apply a green/red colour scale across the numeric body of a pivot."""
        n_rows, n_cols = frame.shape
        if n_cols < 2 or n_rows < 1:
            return
        first = get_column_letter(2)
        last = get_column_letter(n_cols)
        span = f"{first}{header_row + 1}:{last}{header_row + n_rows}"
        good, bad = ("63BE7B", "F8696B") if lower_is_better else ("F8696B", "63BE7B")
        worksheet.conditional_formatting.add(
            span,
            ColorScaleRule(start_type="min", start_color=good,
                           end_type="max", end_color=bad),
        )

    # ----------------------------------------------------------------- readme

    def _readme_lines(self, tables: Dict[str, pd.DataFrame]) -> List[Tuple[str, str]]:
        cfg = self.cfg
        return [
            ("title", "P-Dependent Fusion Rules - 5-Fold CV Evaluation on TRAIN Folds"),
            ("note",  "Generated by evaluate_pdep_train_5fold.py"),
            ("blank", ""),
            ("head",  "WHAT THIS REPORT CONTAINS"),
            ("body",  "Every parametrized fusion rule is evaluated at every value of p, on every "
                      "fusion system, within each cross-validation TRAIN fold. This is the tuning "
                      "stage of a train/test protocol: the held-out test folds are NOT used here."),
            ("blank", ""),
            ("head",  "PROTOCOL"),
            ("body",  f"Subjects are shuffled with seed={cfg.split.seed} and split into "
                      f"{cfg.split.n_folds} folds. Each fold is held out in turn, leaving the "
                      f"remaining folds as the TRAIN subset."),
            ("body",  "Min-max normalization parameters are fitted on the TRAIN submatrix only, "
                      "so held-out subjects never influence the scale their own scores map to."),
            ("body",  "Genuine comparisons are the matrix diagonal; impostor comparisons are all "
                      "off-diagonal entries. Invalid entries are excluded throughout."),
            ("body",  f"Rank-1 identification is measured at FAR = {cfg.metrics.far_target:.0%}: a probe "
                      "counts only if its top match is correct AND clears the FAR threshold."),
            ("blank", ""),
            ("head",  "HOW STATISTICS ARE COMPUTED"),
            ("body",  "Every (system, rule, p) configuration is evaluated in all folds and "
                      "averaged FIRST. Only then is the best p chosen, and then the best rule -- "
                      "both on the MEAN across folds."),
            ("body",  "This ordering matters: choosing p separately inside each fold and then "
                      "averaging would blend several different configurations into one number, "
                      "which is optimistically biased and matches no deployable system. Here "
                      "every reported winner is a single fixed (rule, p) pair applied identically "
                      "in all folds."),
            ("body",  f"Reported for each configuration: mean, std, min, max and a "
                      f"{cfg.metrics.confidence_level:.0%} confidence interval "
                      f"(t-distribution, df = {cfg.split.n_folds - 1}), for both EER and "
                      "Rank-1. The std column shows how stable a configuration is across folds; "
                      "a large std means the mean should be trusted less."),
            ("blank", ""),
            ("head",  "SHEET GUIDE"),
        ] + [
            ("body", f"  {name} - {self.SHEET_TITLES.get(name, '')}")
            for name in tables
        ]

    def _write_readme(self, workbook, tables: Dict[str, pd.DataFrame], styles: dict) -> None:
        worksheet = workbook.active
        worksheet.title = "README"
        worksheet.column_dimensions["A"].width = 105

        fonts = {
            "title": styles["title_font"],
            "head":  Font(name=self.font, bold=True, size=11),
            "note":  styles["note_font"],
            "body":  styles["cell_font"],
            "blank": styles["cell_font"],
        }
        for row, (kind, text) in enumerate(self._readme_lines(tables), start=1):
            cell = worksheet.cell(row=row, column=1, value=text)
            cell.font = fonts[kind]
            cell.alignment = Alignment(wrap_text=False, vertical="top")

    # ------------------------------------------------------------------ write

    def write(self, tables: Dict[str, pd.DataFrame], path: Path) -> Path:
        if not OPENPYXL_AVAILABLE:
            raise RuntimeError("openpyxl is required to write Excel reports. "
                               "Install it, or run with --no-excel.")

        styles = self._styles()
        workbook = Workbook()
        self._write_readme(workbook, tables, styles)

        heatmaps = {
            "Pivot_Mean_EER": True,       # lower is better
            "Pivot_Std_EER": True,        # lower (more stable) is better
            "Pivot_Mean_Rank1": False,    # higher is better
            "Pivot_Std_Rank1": True,      # lower (more stable) is better
        }

        for name, frame in tables.items():
            worksheet = workbook.create_sheet(name[:31])  # Excel caps sheet names
            title = self.SHEET_TITLES.get(name)
            note = f"{len(frame)} rows" if len(frame) > 20 else None
            self._write_table(worksheet, frame, styles, title, note)

            if name in heatmaps:
                header_row = 4 if title else 2
                self._add_heatmap(worksheet, frame, header_row, heatmaps[name])

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        self.log.info("Excel report written to %s", path)
        return path


# ==============================================================================
#  SECTION 9 -- COMMAND-LINE INTERFACE
# ==============================================================================

def parse_sentinel(text: str) -> float:
    """
    Parse the --invalid-value argument, accepting 'nan' alongside numbers.

    argparse's plain `type=float` would accept 'nan' too, but this wrapper
    gives a clearer error message for malformed input.
    """
    cleaned = text.strip().lower()
    if cleaned in {"nan", "none", "null"}:
        return float("nan")
    try:
        return float(cleaned)
    except ValueError as exc:
        raise argparse.ArgumentTypeError(
            f"Invalid sentinel '{text}'. Expected a number or 'nan'."
        ) from exc


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="5-fold CV evaluation of p-dependent fusion rules on TRAIN folds.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    defaults = Config()

    data = parser.add_argument_group("data")
    data.add_argument("--mat-file", default=defaults.data.mat_file,
                      help="Path to the .mat file containing the score matrices.")
    data.add_argument("--invalid-value", type=parse_sentinel,
                      default=defaults.data.invalid_value,
                      help="Sentinel marking invalid/missing scores. "
                           "Accepts a number or the literal 'nan'.")
    data.add_argument("--no-normalize", action="store_true",
                      help="Skip min-max normalization. Use when the score "
                           "matrices are already normalized to [0, 1].")
    data.add_argument("--skip-range-check", action="store_true",
                      help="With --no-normalize, do not verify that scores "
                           "lie within [0, 1].")

    split = parser.add_argument_group("cross-validation")
    split.add_argument("--n-folds", type=int, default=defaults.split.n_folds,
                       help="Number of cross-validation folds.")
    split.add_argument("--seed", type=int, default=defaults.split.seed,
                       help="Random seed controlling the subject split.")
    split.add_argument("--no-shuffle", action="store_true",
                       help="Split subjects in their original order.")

    sweep = parser.add_argument_group("sweep")
    sweep.add_argument("--p-min", type=float, default=defaults.sweep.p_min)
    sweep.add_argument("--p-max", type=float, default=defaults.sweep.p_max)
    sweep.add_argument("--p-step", type=float, default=defaults.sweep.p_step,
                       help="Granularity of the p sweep.")
    sweep.add_argument("--rules", nargs="+", default=None,
                       help=f"Subset of rules to evaluate. Choices: {sorted(FUSION_RULES)}")
    sweep.add_argument("--systems", nargs="+", default=None,
                       help=f"Subset of systems to evaluate. Choices: {sorted(SYSTEM_DEFINITIONS)}")

    metrics = parser.add_argument_group("metrics")
    metrics.add_argument("--far-target", type=float, default=defaults.metrics.far_target,
                         help="FAR at which Rank-1 identification is measured.")
    metrics.add_argument("--confidence-level", type=float,
                         default=defaults.metrics.confidence_level,
                         help="Confidence level for the reported interval.")

    output = parser.add_argument_group("output")
    output.add_argument("--output", default=defaults.output.output_xlsx,
                        help="Filename of the Excel report.")
    output.add_argument("--output-dir", default=defaults.output.output_dir,
                        help="Directory to write outputs into.")
    output.add_argument("--no-excel", action="store_true", help="Skip the .xlsx report.")
    output.add_argument("--no-csv", action="store_true", help="Skip the .csv exports.")
    output.add_argument("--quiet", action="store_true", help="Reduce console output.")

    return parser


def config_from_args(args: argparse.Namespace) -> Config:
    """Translate parsed CLI arguments into a Config object."""
    return Config(
        data=DataConfig(
            mat_file=args.mat_file,
            invalid_value=args.invalid_value,
            normalize=not args.no_normalize,
            validate_range_if_not_normalized=not args.skip_range_check,
        ),
        split=SplitConfig(
            n_folds=args.n_folds,
            seed=args.seed,
            shuffle=not args.no_shuffle,
        ),
        sweep=SweepConfig(
            p_min=args.p_min,
            p_max=args.p_max,
            p_step=args.p_step,
            rules=args.rules,
            systems=args.systems,
        ),
        metrics=MetricConfig(
            far_target=args.far_target,
            confidence_level=args.confidence_level,
        ),
        output=OutputConfig(
            output_xlsx=args.output,
            output_dir=args.output_dir,
            write_excel=not args.no_excel,
            write_csv=not args.no_csv,
        ),
        verbose=not args.quiet,
    )


def export_csvs(tables: Dict[str, pd.DataFrame], directory: Path) -> None:
    """Write each summary table alongside the workbook as a .csv."""
    csv_dir = directory / "csv"
    csv_dir.mkdir(parents=True, exist_ok=True)
    for name, frame in tables.items():
        frame.to_csv(csv_dir / f"{name}.csv", index=False)
    logging.getLogger("export").info("CSV exports written to %s", csv_dir)


def print_console_summary(tables: Dict[str, pd.DataFrame]) -> None:
    """Show the headline result without needing to open the workbook."""
    best = tables.get("Overall_Best_per_System")
    if best is None or best.empty:
        return

    print("\n" + "=" * 86)
    print("WINNING RULE PER SYSTEM  (lowest MEAN TRAIN EER across folds)")
    print("=" * 86)
    print(f"{'System':<8}{'Rule':<14}{'p':>5}"
          f"{'EER mean':>11}{'EER std':>10}{'±CI95':>9}"
          f"{'R1 mean':>10}{'R1 std':>9}")
    print("-" * 86)
    for row in best.itertuples(index=False):
        print(f"{row.System:<8}{str(row.Rule):<14}{row.p_selected:>5.2f}"
              f"{row.Train_EER_mean * 100:>10.2f}%{row.Train_EER_std * 100:>9.2f}%"
              f"{row.Train_EER_CI95 * 100:>8.2f}%"
              f"{row.Train_Rank1_mean * 100:>9.2f}%{row.Train_Rank1_std * 100:>8.2f}%")
    print("=" * 86 + "\n")


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = build_arg_parser().parse_args(argv)
    config = config_from_args(args)

    logging.basicConfig(
        level=logging.INFO if config.verbose else logging.WARNING,
        format="%(asctime)s  %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
        stream=sys.stdout,
    )
    log = logging.getLogger("main")

    try:
        evaluator = PDependentTrainEvaluator(config)
        evaluator.load_data()
        evaluator.run()
        tables = evaluator.summarize()
    except (FileNotFoundError, KeyError, ValueError, RuntimeError) as exc:
        log.error("%s: %s", type(exc).__name__, exc)
        return 1

    out_dir = Path(config.output.output_dir)

    if config.output.write_csv:
        export_csvs(tables, out_dir)

    if config.output.write_excel:
        try:
            ExcelReportWriter(config).write(tables, out_dir / config.output.output_xlsx)
        except RuntimeError as exc:
            log.error("%s", exc)
            return 1

    print_console_summary(tables)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
